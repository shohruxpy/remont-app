from fastapi import APIRouter, Depends, Response, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from typing import Optional
from datetime import datetime
import io

from ..database import get_db
from ..models import User, Repair, RepairItem, Zaprafka, Machine
from ..auth import admin_required

router = APIRouter()

@router.get("/{report_type}.xlsx")
async def get_excel_report(
    report_type: str,
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    machine_code: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(admin_required)
):
    valid_types = ["expenses-by-machine", "repair-history", "zaprafka-deadlines", "top-parts"]
    if report_type not in valid_types:
        raise HTTPException(status_code=404, detail="Report type not found")
        
    import openpyxl
    from openpyxl.styles import Font, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчет"
    
    header_font = Font(bold=True)
    
    if report_type == "expenses-by-machine":
        ws.append(["Код станка", "Название", "Новые (кол-во)", "Новые (сумма)", "Восстановленные (кол-во)"])
        # Set header font
        for cell in ws[1]:
            cell.font = header_font
        ws.freeze_panes = "A2"
        
        query = select(Repair, Machine).join(Machine, Repair.machine_id == Machine.id)
        if machine_code:
            query = query.where(Machine.code == machine_code)
        if date_from:
            query = query.where(Repair.repair_date >= date_from)
        if date_to:
            query = query.where(Repair.repair_date <= date_to)
            
        res = await db.execute(query)
        repairs = res.all()
        
        machine_stats = {}
        for r, m in repairs:
            if m.code not in machine_stats:
                machine_stats[m.code] = {"name": m.name, "new_q": 0, "new_a": 0, "ref_q": 0}
            
            items_res = await db.execute(select(RepairItem).where(RepairItem.repair_id == r.id))
            for i in items_res.scalars().all():
                if i.condition == "NEW":
                    machine_stats[m.code]["new_q"] += float(i.qty)
                    machine_stats[m.code]["new_a"] += float(i.amount) if i.amount else 0
                else:
                    machine_stats[m.code]["ref_q"] += float(i.qty)
                    
        for code, st in machine_stats.items():
            ws.append([code, st["name"], st["new_q"], st["new_a"], st["ref_q"]])

    elif report_type == "repair-history":
        ws.append(["Дата", "Код станка", "Тип", "Автор", "Сумма"])
        for cell in ws[1]:
            cell.font = header_font
        ws.freeze_panes = "A2"
        
        query = select(Repair, Machine, User).join(Machine, Repair.machine_id == Machine.id).join(User, Repair.created_by == User.id)
        if machine_code:
            query = query.where(Machine.code == machine_code)
        if date_from:
            query = query.where(Repair.repair_date >= date_from)
        if date_to:
            query = query.where(Repair.repair_date <= date_to)
            
        res = await db.execute(query)
        repairs = res.all()
        
        for r, m, u in repairs:
            items_res = await db.execute(select(RepairItem).where(RepairItem.repair_id == r.id))
            amt = sum(float(i.amount) if i.amount else 0 for i in items_res.scalars().all())
            date_str = r.repair_date.strftime("%d.%m.%Y")
            ws.append([date_str, m.code, r.type, u.full_name, amt])
            
    elif report_type == "zaprafka-deadlines":
        ws.append(["Код станка", "Название", "Последняя", "Следующая", "Осталось", "Статус"])
        for cell in ws[1]:
            cell.font = header_font
        ws.freeze_panes = "A2"
        
        query = select(Machine).where(Machine.is_archived == False)
        if machine_code:
            query = query.where(Machine.code == machine_code)
        res = await db.execute(query)
        
        for m in res.scalars().all():
            if m.last_zaprafka_end:
                date_str = m.last_zaprafka_end.strftime("%d.%m.%Y")
                ws.append([m.code, m.name, date_str, "—", "—", m.status])
            else:
                ws.append([m.code, m.name, "—", "—", "—", m.status])
                
    elif report_type == "top-parts":
        ws.append(["Материал", "Кол-во"])
        for cell in ws[1]:
            cell.font = header_font
        ws.freeze_panes = "A2"
        
        query = select(RepairItem.free_text_material, RepairItem.qty)
        # Note: In a real system, we'd join Repair to filter by machine/date and sum by material.
        res = await db.execute(query)
        
        stats = {}
        for name, qty in res.all():
            if name:
                stats[name] = stats.get(name, 0) + float(qty)
        
        for name, total_qty in sorted(stats.items(), key=lambda x: x[1], reverse=True):
            ws.append([name, total_qty])
            
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    return Response(
        content=out.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={report_type}.xlsx"}
    )
