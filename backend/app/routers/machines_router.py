from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from typing import List, Optional
from pydantic import BaseModel
import uuid
import io

from ..database import get_db
from ..models import Machine, AuditLog, User
from ..schemas import MachineOut, MachineCreate, MachineUpdate
from ..auth import get_current_user, admin_required

router = APIRouter()

@router.get("/", response_model=List[MachineOut])
async def read_machines(db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user), search: Optional[str] = None, status: Optional[str] = None, page: int = 1, page_size: Optional[int] = 100):
    query = select(Machine).where(Machine.is_archived == False).order_by(Machine.code.asc())
    if search:
        query = query.where(Machine.code.ilike(f"%{search}%") | Machine.name.ilike(f"%{search}%"))
    if status:
        query = query.where(Machine.status == status)
    
    if page_size and page_size > 0:
        query = query.offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(query)
    return result.scalars().all()

@router.get("/{code}", response_model=MachineOut)
async def read_machine(code: str, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    result = await db.execute(select(Machine).where(Machine.code == code))
    machine = result.scalars().first()
    if not machine:
        raise HTTPException(status_code=404, detail="Machine not found")
    return machine

@router.post("/", response_model=MachineOut)
async def create_machine(machine: MachineCreate, db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    result = await db.execute(select(Machine).where(Machine.code == machine.code))
    if result.scalars().first():
        raise HTTPException(status_code=400, detail="Machine with this code already exists")
        
    m_dict = machine.model_dump()
    if m_dict.get("last_zaprafka_end") and hasattr(m_dict["last_zaprafka_end"], "tzinfo") and m_dict["last_zaprafka_end"].tzinfo:
        m_dict["last_zaprafka_end"] = m_dict["last_zaprafka_end"].replace(tzinfo=None)
    db_machine = Machine(**m_dict, created_by=current_user.id)
    db.add(db_machine)
    await db.flush()
    
    log = AuditLog(user_id=current_user.id, action="CREATE", entity="MACHINE", entity_id=db_machine.id, after=m_dict)
    db.add(log)
    await db.commit()
    await db.refresh(db_machine)
    return db_machine

@router.put("/{machine_id}", response_model=MachineOut)
async def update_machine(machine_id: uuid.UUID, machine_data: MachineUpdate, db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    result = await db.execute(select(Machine).where(Machine.id == machine_id))
    machine = result.scalars().first()
    if not machine:
        raise HTTPException(status_code=404, detail="Machine not found")
        
    update_dict = machine_data.model_dump(exclude_unset=True)
    if update_dict.get("last_zaprafka_end") and hasattr(update_dict["last_zaprafka_end"], "tzinfo") and update_dict["last_zaprafka_end"].tzinfo:
        update_dict["last_zaprafka_end"] = update_dict["last_zaprafka_end"].replace(tzinfo=None)
    for field, val in update_dict.items():
        setattr(machine, field, val)
        
    log = AuditLog(user_id=current_user.id, action="UPDATE", entity="MACHINE", entity_id=machine.id, after=update_dict)
    db.add(log)
    await db.commit()
    await db.refresh(machine)
    return machine

@router.delete("/{machine_id}")
async def delete_machine(machine_id: uuid.UUID, db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    result = await db.execute(select(Machine).where(Machine.id == machine_id))
    machine = result.scalars().first()
    if not machine:
        raise HTTPException(status_code=404, detail="Machine not found")
        
    machine.is_archived = True
    log = AuditLog(user_id=current_user.id, action="DELETE", entity="MACHINE", entity_id=machine.id)
    db.add(log)
    await db.commit()
    return {"status": "success", "message": "Machine deleted"}


@router.post("/import")
async def import_machines(dry_run: bool = False, file: UploadFile = File(...), db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    import openpyxl
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheet = wb.active
    
    results = []
    errors = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]: continue
        code, name, co_order, cost_center = row[0:4]
        code = str(code)
        
        # Check if code exists
        existing = await db.execute(select(Machine).where(Machine.code == code))
        if existing.scalars().first():
            errors.append(f"Row {row_idx}: Machine {code} already exists")
            continue
            
        machine_data = MachineCreate(
            code=code, name=str(name) if name else f"Machine {code}",
            sap_co_order=str(co_order) if co_order else None,
            sap_cost_center=str(cost_center) if cost_center else None
        )
        
        if not dry_run:
            db_machine = Machine(**machine_data.model_dump(), created_by=current_user.id)
            db.add(db_machine)
            await db.flush()
            log = AuditLog(user_id=current_user.id, action="CREATE", entity="MACHINE", entity_id=db_machine.id, after=machine_data.model_dump())
            db.add(log)
        results.append(f"Row {row_idx}: Added {code}")
        
    if not dry_run:
        await db.commit()
        
    return {"results": results, "errors": errors}

class QrSheetRequest(BaseModel):
    codes: List[str]


@router.post("/qr-sheet")
async def generate_qr_sheet(req: QrSheetRequest, db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    from fpdf import FPDF
    import os
    pdf = FPDF()
    pdf.add_page()
    font_path = os.path.join(os.path.dirname(__file__), "../../../Roboto-Regular.ttf")
    if os.path.exists(font_path):
        pdf.add_font('Roboto', '', font_path, uni=True)
        pdf.set_font('Roboto', '', 12)
    else:
        pdf.set_font("Arial", size=12)
        
    for code in req.codes:
        res = await db.execute(select(Machine).where(Machine.code == code))
        m = res.scalars().first()
        name = m.name if m else ""
        pdf.cell(200, 10, txt=f"EQ:{code} - {name}", ln=True, align='C')
    
    return Response(content=bytes(pdf.output()), media_type="application/pdf")

@router.get("/{code}/qr")
async def generate_qr_single(code: str, db: AsyncSession = Depends(get_db), current_user: User = Depends(admin_required)):
    from fpdf import FPDF
    import os
    pdf = FPDF()
    pdf.add_page()
    font_path = os.path.join(os.path.dirname(__file__), "../../../Roboto-Regular.ttf")
    if os.path.exists(font_path):
        pdf.add_font('Roboto', '', font_path, uni=True)
        pdf.set_font('Roboto', '', 12)
    else:
        pdf.set_font("Arial", size=12)
        
    res = await db.execute(select(Machine).where(Machine.code == code))
    m = res.scalars().first()
    name = m.name if m else ""
    pdf.cell(200, 10, txt=f"EQ:{code} - {name}", ln=True, align='C')
    
    return Response(content=bytes(pdf.output()), media_type="application/pdf")
